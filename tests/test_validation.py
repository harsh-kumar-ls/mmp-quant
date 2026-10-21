"""Validation harness (build order step 4).

TEST DATA ONLY. Every "manual score" in this file is synthetic, constructed so
the correct answer is known in advance (perfect agreement -> 1, reversed -> -1,
...). They test the harness's arithmetic and bookkeeping, say nothing about
whether the system is right, and exist only in temporary directories -- never
in data/.
"""

from __future__ import annotations

import json
import math

import pandas as pd
import pytest

from mmpquant.validation import stats
from mmpquant.validation.ground_truth import GroundTruthError, load_ground_truth
from mmpquant.validation.harness import Settings, ValidationError, run_validation

CFG = {"scoring": {"principles": ["redundancy", "coherence"]}}
FAST = Settings(n_resamples=400, seed=0)


def make_run(runs, vid, score, *, assessed=None, segs=None, coh=None, coh_alt=None,
             config=CFG):
    """A minimal result.json with only the fields the harness reads."""
    segs = segs if segs is not None else [score, score]
    payload = {
        "video_id": vid,
        "result": {"score": score, "status": "OK" if score is not None else "N/A"},
        "assessed_score": assessed if assessed is not None else score,
        "segments": [{"index": i, "score": s, "assessed_score": s} for i, s in enumerate(segs)],
        "config": config,
    }
    if coh is not None:
        payload["coherence"] = {
            "result": {"score": coh, "status": "OK"},
            "alternative": {"score": coh_alt if coh_alt is not None else coh,
                            "weights": {"background_music": 0.5, "off_topic": 0.5}},
            "segments": [{"index": i, "score": coh,
                          "parts": {"background_music": coh / 100, "off_topic": 0.0}}
                         for i in range(len(segs))],
        }
    (runs / vid).mkdir(parents=True)
    (runs / vid / "result.json").write_text(json.dumps(payload))


def write_gt(path, rows, header="video_id,manual_score", comment="# TEST DATA - synthetic"):
    path.write_text("\n".join([comment, header, *rows]) + "\n")
    return path


@pytest.fixture
def corpus(tmp_path):
    runs = tmp_path / "runs"
    system = [0, 5, 10, 20, 30, 40, 50, 60, 70, 80]
    for i, s in enumerate(system):
        make_run(runs, f"v{i}", float(s), coh=float(s), coh_alt=float(100 - s))
    return runs, system


def validate(tmp_path, gt_path, runs, principle="redundancy", config=CFG, settings=FAST):
    return run_validation(principle, load_ground_truth(gt_path), runs, tmp_path / "out",
                          current_config=config, settings=settings)


class TestKnownAnswers:
    def test_identical_order_gives_one(self, tmp_path, corpus):
        runs, system = corpus
        gt = write_gt(tmp_path / "m.csv", [f"v{i},{s}" for i, s in enumerate(system)])
        r = validate(tmp_path, gt, runs)
        assert r["held_out"]["spearman"] == 1.0
        assert r["nested_loo"]["spearman"] == 1.0

    def test_reversed_order_gives_minus_one(self, tmp_path, corpus):
        runs, system = corpus
        gt = write_gt(tmp_path / "m.csv", [f"v{i},{100 - s}" for i, s in enumerate(system)])
        r = validate(tmp_path, gt, runs)
        assert r["held_out"]["spearman"] == -1.0

    def test_the_professors_scale_does_not_matter(self, tmp_path, corpus):
        runs, system = corpus
        likert = [1, 1, 2, 2, 3, 3, 4, 4, 5, 5]   # a 1-5 scale, same order
        a = validate(tmp_path, write_gt(tmp_path / "a.csv",
                                        [f"v{i},{s}" for i, s in enumerate(system)]), runs)
        b = validate(tmp_path, write_gt(tmp_path / "b.csv",
                                        [f"v{i},{x * 20}" for i, x in enumerate(likert)]), runs)
        c = validate(tmp_path, write_gt(tmp_path / "c.csv",
                                        [f"v{i},{x}" for i, x in enumerate(likert)]), runs)
        assert b["exploratory_all_videos"]["whole_video"]["spearman"] == \
            c["exploratory_all_videos"]["whole_video"]["spearman"]
        assert a["exploratory_all_videos"]["whole_video"]["spearman"] == 1.0

    def test_outputs_are_written(self, tmp_path, corpus):
        runs, system = corpus
        validate(tmp_path, write_gt(tmp_path / "m.csv",
                                    [f"v{i},{s}" for i, s in enumerate(system)]), runs)
        for name in ("validation.json", "report.md", "residuals.csv", "scatter.png"):
            assert (tmp_path / "out" / name).exists()


class TestNothingIsSilentlyDropped:
    def test_missing_runs_and_na_scores_are_listed_not_zeroed(self, tmp_path, corpus):
        runs, system = corpus
        make_run(runs, "silent", None)
        rows = [f"v{i},{s}" for i, s in enumerate(system)] + ["silent,3", "ghost,50"]
        r = validate(tmp_path, write_gt(tmp_path / "m.csv", rows), runs)
        assert r["excluded"]["system_na"] == [{"video_id": "silent"}]
        assert r["excluded"]["no_system_score"] == [{"video_id": "ghost"}]
        assert "no run" in r["excluded"]["not_scored_reason"]["ghost"]
        assert r["exploratory_all_videos"]["whole_video"]["n_videos"] == 10

    def test_blank_manual_cells_are_not_rated_not_zero(self, tmp_path, corpus):
        _runs, system = corpus
        rows = [f"v{i},{s}" for i, s in enumerate(system)] + ["v_extra,"]
        gt = load_ground_truth(write_gt(tmp_path / "m.csv", rows))
        assert gt.n_unrated_rows == 1 and "v_extra" not in set(gt.table["video_id"])

    def test_runs_scored_under_another_config_are_refused(self, tmp_path, corpus):
        runs, system = corpus
        gt = write_gt(tmp_path / "m.csv", [f"v{i},{s}" for i, s in enumerate(system)])
        with pytest.raises(ValidationError, match="different config"):
            validate(tmp_path, gt, runs, config={"something": "else"})

    def test_too_few_videos_is_refused(self, tmp_path, corpus):
        runs, _ = corpus
        with pytest.raises(ValidationError, match="at least"):
            validate(tmp_path, write_gt(tmp_path / "m.csv", ["v1,1", "v2,2", "v3,3"]), runs)


class TestGroundTruthFile:
    def test_non_numeric_score_is_an_error(self, tmp_path):
        with pytest.raises(GroundTruthError, match="non-numeric"):
            load_ground_truth(write_gt(tmp_path / "m.csv", ["v1,high"]))

    def test_duplicate_rows_are_an_error(self, tmp_path):
        with pytest.raises(GroundTruthError, match="duplicate"):
            load_ground_truth(write_gt(tmp_path / "m.csv", ["v1,1", "v1,2"]))

    def test_a_system_output_file_is_refused(self, tmp_path):
        with pytest.raises(GroundTruthError, match="system output"):
            load_ground_truth(write_gt(tmp_path / "m.csv", ["v1,1"],
                                       header="video_id,system_score"))

    def test_dummy_marker_is_detected_and_flagged(self, tmp_path, corpus):
        runs, system = corpus
        r = validate(tmp_path, write_gt(tmp_path / "m.csv",
                                        [f"v{i},{s}" for i, s in enumerate(system)]), runs)
        assert r["is_dummy_ground_truth"] is True
        assert "TEST DATA — NOT A RESULT" in (tmp_path / "out" / "report.md").read_text()

    def test_unmarked_file_is_not_flagged(self, tmp_path):
        gt = load_ground_truth(write_gt(tmp_path / "m.csv", ["v1,1"], comment=""))
        assert gt.is_dummy is False

    def test_two_raters_mean_and_agreement(self, tmp_path, corpus):
        runs, system = corpus
        rows = [f"v{i},{s},{s + 1}" for i, s in enumerate(system)]
        gt_path = write_gt(tmp_path / "m.csv", rows, header="video_id,rater_1_score,rater_2_score")
        gt = load_ground_truth(gt_path)
        assert gt.raters == ["rater_1_score", "rater_2_score"]
        assert gt.table.loc[0, "manual"] == pytest.approx(0.5)
        r = validate(tmp_path, gt_path, runs)
        assert r["inter_rater"][0]["spearman"] == 1.0

    def test_a_folder_of_filled_annotation_sheets(self, tmp_path, corpus):
        runs, system = corpus
        sheets = tmp_path / "sheets"
        sheets.mkdir()
        header = "video_id,segment_index,start_s,end_s,timecode,rater_1_score,rater_2_score"
        for i, s in enumerate(system):
            write_gt(sheets / f"v{i}_sheet.csv",
                     [f"v{i},0,0,30,0:00-0:30,{s},{s}", f"v{i},1,30,60,0:30-1:00,{s},"],
                     header=header)
        gt = load_ground_truth(sheets)
        assert gt.level == "segment" and len(gt.table) == 20
        r = validate(tmp_path, sheets, runs)
        assert r["level"] == "segment" and r["exploratory_all_videos"]["whole_video"]["n_rows"] == 20


class TestHonestSplitting:
    def test_split_is_by_video_and_reproducible(self):
        ids = [f"v{i}" for i in range(10)]
        a = stats.split_videos(ids, 0.6, seed=0)
        assert a == stats.split_videos(ids, 0.6, seed=0)
        assert not set(a[0]) & set(a[1]) and len(a[0]) == 6

    def test_segments_of_one_video_never_straddle_the_split(self, tmp_path, corpus):
        runs, system = corpus
        rows = [f"v{i},{j},{s}" for i, s in enumerate(system) for j in (0, 1)]
        r = validate(tmp_path, write_gt(tmp_path / "m.csv", rows,
                                        header="video_id,segment_index,manual_score"), runs)
        res = pd.read_csv(tmp_path / "out" / "residuals.csv")
        assert (res.groupby("video_id")["split"].nunique() == 1).all()
        assert set(r["split"]["tune_videos"]).isdisjoint(r["split"]["held_out_videos"])

    def test_variant_is_chosen_on_the_tune_videos_only(self, tmp_path, corpus):
        runs, system = corpus
        # coherence: music_only follows `system`, music_plus_off_topic is reversed.
        # Manual scores agree with the reversed one -> it must be chosen on tune.
        gt = write_gt(tmp_path / "m.csv", [f"v{i},{100 - s}" for i, s in enumerate(system)])
        r = validate(tmp_path, gt, runs, principle="coherence")
        assert r["tune"]["chosen_variant"] == "music_plus_off_topic"
        assert r["held_out"]["spearman"] == 1.0
        assert r["held_out"]["primary_for_reference"]["spearman"] == -1.0

    def test_cluster_bootstrap_resamples_videos(self):
        x = [1, 1, 2, 2, 3, 3, 4, 4]
        y = [1, 1, 2, 2, 3, 3, 4, 4]
        lo, hi, n = stats.bootstrap_ci(x, y, groups=list("aabbccdd"), n_resamples=200)
        assert n > 0 and lo <= hi <= 1.0

    def test_undefined_correlation_is_nan_not_zero(self):
        assert math.isnan(stats.spearman([0, 0, 0, 0], [1, 2, 3, 4]))


class TestCommandLine:
    def test_refuses_without_manual_scores(self, tmp_path, monkeypatch):
        from typer.testing import CliRunner

        from mmpquant.cli import app

        result = CliRunner().invoke(app, ["validate", "-g", str(tmp_path / "none.csv")])
        assert result.exit_code == 1
        assert "Nothing is computed from placeholder data" in result.output


class TestRatingWorkbooks:
    """20 Oct 2026: one Excel workbook per rater (ratings_<name>.xlsx)."""

    def _workbook(self, path, scores, *, overall=None):
        """Build a blank workbook, then fill it with SYNTHETIC test scores."""
        from openpyxl import load_workbook

        from mmpquant.rating_workbook import build_workbook

        build_workbook([(f"v{i}", 60.0) for i in range(len(scores))], path, segment_length_s=30)
        wb = load_workbook(path)
        ws = wb["Ratings"]
        for row in ws.iter_rows(min_row=2):
            vid = int(row[0].value[1:])
            row[5].value = scores[vid]          # redundancy_score
            row[6].value = 5 - scores[vid] // 25 if scores[vid] is not None else None
        if overall:
            for row in wb["Video overall"].iter_rows(min_row=2):
                row[2].value = overall[int(row[0].value[1:])]
        wb.save(path)
        return path

    def test_blank_workbook_has_no_system_scores_and_every_segment(self, tmp_path):
        from openpyxl import load_workbook

        from mmpquant.rating_workbook import build_workbook

        p = build_workbook([("a", 65.0), ("b", 30.0)], tmp_path / "t.xlsx", segment_length_s=30)
        wb = load_workbook(p)
        assert wb.sheetnames == ["Read me", "Ratings", "Video overall"]
        rows = list(wb["Ratings"].iter_rows(min_row=2, values_only=True))
        assert [(r[0], r[1]) for r in rows] == [("a", 0), ("a", 1), ("a", 2), ("b", 0)]
        headers = [c.value for c in wb["Ratings"][1]]
        assert not any("system" in str(h) for h in headers)

    def test_two_raters_round_trip_through_validation(self, tmp_path, corpus):
        runs, system = corpus
        folder = tmp_path / "ratings"
        folder.mkdir()
        self._workbook(folder / "ratings_alice.xlsx", system)
        self._workbook(folder / "ratings_bob.xlsx", [s + 1 for s in system])
        gt = load_ground_truth(folder, principle="redundancy")
        assert gt.raters == ["alice_score", "bob_score"] and gt.level == "segment"
        r = run_validation("redundancy", gt, runs, tmp_path / "out",
                           current_config=CFG, settings=FAST)
        assert r["exploratory_all_videos"]["whole_video"]["spearman"] == 1.0
        assert r["inter_rater"][0]["spearman"] == 1.0

    def test_video_overall_sheet_gives_video_level(self, tmp_path, corpus):
        _runs, system = corpus
        folder = tmp_path / "ratings"
        folder.mkdir()
        self._workbook(folder / "ratings_alice.xlsx", system, overall=system)
        gt = load_ground_truth(folder, principle="redundancy", level="video")
        assert gt.level == "video" and len(gt.table) == len(system)

    def test_workbooks_need_a_principle(self, tmp_path, corpus):
        _runs, system = corpus
        folder = tmp_path / "ratings"
        folder.mkdir()
        self._workbook(folder / "ratings_alice.xlsx", system)
        with pytest.raises(GroundTruthError, match="say which one"):
            load_ground_truth(folder)
